"""Statistical readers use validated snapshots and never create a GUI."""
import json
from pathlib import Path
import tempfile
import unittest

from flexdmo_app.core import save_frames
from flexdmo_app.tests import test_core
from results_output.MIGD_table.main import run
from utils.result_io import load_result_from_files


class ResultReaderTests(unittest.TestCase):
    def setUp(self):
        fixture = test_core.CoreTests()
        fixture.setUp()
        self.frame = fixture.frame
        self.directory = tempfile.TemporaryDirectory(prefix="FlexDMO-统计-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.path = self.root / "子目录" / "结果.json"
        save_frames(self.path, [self.frame])

    def test_single_path_and_overlapping_inputs_are_supported(self):
        self.assertEqual(len(list(load_result_from_files(str(self.path)))), 1)
        results = list(load_result_from_files([self.path, self.root, self.path.parent]))
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["file_path"], str(self.path))
        self.assertEqual(sorted(results[0]["runtime_populations"]), [1])

    def test_invalid_json_and_missing_path_do_not_hide_valid_result(self):
        invalid = self.root / "invalid.json"
        invalid.write_text("{invalid", encoding="utf-8")
        errors = []
        results = list(load_result_from_files([invalid, self.root / "missing", self.path],
                       on_error=lambda path, error: errors.append((path, error))))
        self.assertEqual(len(results), 1)
        self.assertEqual(len(errors), 2)
        self.assertTrue(all(isinstance(error, (ValueError, OSError)) for _, error in errors))

    def test_invalid_schema_is_reported_without_aborting_the_scan(self):
        malformed = [
            {"settings": []},
            {"settings": dict(self.frame["settings"], problem_params=[])},
            {"settings": self.frame["settings"], "information": {"0": []}},
            {"settings": self.frame["settings"], "information": {"0": {"1": None}}},
        ]
        errors = []
        for index, data in enumerate(malformed):
            (self.root / f"bad-{index}.json").write_text(json.dumps(data), encoding="utf-8")
        results = list(load_result_from_files(self.root,
                       on_error=lambda path, error: errors.append((path, error))))
        self.assertEqual(len(results), 1)
        self.assertEqual(len(errors), len(malformed))

    def test_default_warning_names_bad_file(self):
        invalid = self.root / "invalid.json"
        invalid.write_text("null", encoding="utf-8")
        with self.assertWarnsRegex(RuntimeWarning, "invalid.json"):
            self.assertEqual(list(load_result_from_files(invalid)), [])

    def test_directory_scan_accepts_uppercase_json_extension(self):
        uppercase = self.root / "大写.JSON"
        save_frames(uppercase, [self.frame])
        self.assertEqual(len(list(load_result_from_files(self.root))), 2)

    def test_error_callback_exceptions_propagate(self):
        def reject(path, error):
            raise RuntimeError("caller canceled scan")
        with self.assertRaisesRegex(RuntimeError, "caller canceled"):
            list(load_result_from_files(self.root / "missing", on_error=reject))

    def test_real_statistics_export_works_without_result_dialogs(self):
        output = self.root / "report"
        run({"input_paths": [self.path], "output_path": str(output)})
        paths = list(output.glob("*.xlsx"))
        self.assertEqual(len(paths), 1)
        from openpyxl import load_workbook
        book = load_workbook(paths[0], read_only=True)
        try:
            self.assertEqual(book["MIGD"].cell(2, 1).value, "CDP1")
            self.assertIn("NoResponse", book["MIGD"].cell(1, 3).value)
            self.assertIsNotNone(book["MIGD"].cell(2, 3).value)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
