import importlib
import os
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from algorithms.search_algorithm.NSGA2.main import NSGA2
from problems.benchmark.CDP1.main import CDP1
from results_output.MIGD_table.main import find_first_and_second, run as build_migd_table


class DynamicResponseWorkflowTests(unittest.TestCase):
    def test_every_response_strategy_handles_three_environments(self):
        strategy_names = ("DIP", "MDA", "MDP", "NoResponse", "RNN")
        for strategy_name in strategy_names:
            with self.subTest(strategy=strategy_name):
                strategy_module = importlib.import_module(
                    f"algorithms.response_strategy.{strategy_name}.main"
                )
                strategy_class = getattr(strategy_module, strategy_name)
                problem = CDP1(
                    decision_num=5,
                    n=10,
                    tau=1,
                    solution_num=6,
                    total_evaluate_time=3,
                )
                algorithm = NSGA2(seed=11)
                algorithm.optimize(problem, strategy_class())

                self.assertEqual(list(algorithm.history["runtime"]), [0, 1, 2])
                final_population = list(
                    list(algorithm.history["runtime"][2].values())
                )[-1]
                self.assertEqual(final_population.n, problem.solution_num)
                self.assertTrue(problem.is_ended())


class ResultOutputWorkflowTests(unittest.TestCase):
    def test_first_and_second_supports_one_or_missing_values(self):
        self.assertEqual(find_first_and_second([0.4]), (0, None))
        self.assertEqual(find_first_and_second([None, 0.5, 0.2]), (2, 1))
        self.assertEqual(find_first_and_second([]), (None, None))

    def test_migd_table_supports_a_single_algorithm(self):
        # A compact valid result snapshot is enough to exercise workbook output.
        algorithm = NSGA2(seed=3)
        problem = CDP1(3, 10, 1, 4, 1)
        algorithm.optimize(problem, _NoResponse())
        result = {
            "settings": {
                "response_strategy_class": "NoResponse",
                "search_algorithm_class": "NSGA2",
                "problem_class": "CDP1",
                "problem_params": {"n": 10, "tau": 1},
            },
            "runtime_populations": algorithm.history["runtime"],
        }

        with tempfile.TemporaryDirectory() as directory:
            with patch(
                "results_output.MIGD_table.main.load_result_from_files",
                return_value=iter([result]),
            ):
                output_file = build_migd_table(
                    {"input_paths": ["unused.json"], "output_path": directory}
                )
            self.assertTrue(os.path.isfile(output_file))
            workbook = load_workbook(output_file, read_only=True)
            self.assertEqual(workbook.active.cell(1, 3).value, "NoResponse-NSGA2")
            self.assertEqual(workbook.active.cell(3, 3).value, "-")
            workbook.close()


class _NoResponse:
    def response(self, population, problem, algorithm):
        return population


if __name__ == "__main__":
    unittest.main()
