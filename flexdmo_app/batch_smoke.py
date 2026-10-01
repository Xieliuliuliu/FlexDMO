"""Real batch-process, UI, export, replay, failure and cancellation integration."""
import time


def main():
    import main_qt
    from PySide6.QtCore import Qt, QTimer
    from PySide6.QtWidgets import QApplication
    from flexdmo_app.window import FlexDMOWindow
    from flexdmo_app.batch_runner import BatchRunner
    from flexdmo_app.core import ROOT, load_frames
    from openpyxl import load_workbook
    import copy

    app = QApplication([])
    app.setStyle("Fusion")
    window = FlexDMOWindow()
    window.show()
    batch = window.batch
    output = ROOT / "results" / "batch_smoke"
    def wait(predicate, timeout=60):
        deadline = time.monotonic() + timeout
        while not predicate():
            app.processEvents()
            if time.monotonic() >= deadline:
                raise TimeoutError("Batch integration timed out")
            time.sleep(0.01)
        app.processEvents()
    try:
        window.switch_workspace(1)
        for kind, wanted in {"dynamic": {"NoResponse"}, "search": {"NSGAII", "SPEA2"}, "problem": {"CDP1", "CDP6"}}.items():
            choices = batch.lists[kind]
            for i in range(choices.count()):
                item = choices.item(i)
                item.setCheckState(Qt.CheckState.Checked if item.data(Qt.ItemDataRole.UserRole)["name"] in wanted else Qt.CheckState.Unchecked)
        for key, value in {"decision_num": 3, "solution_num": 6, "total_evaluate_time": 3,
                           "repeats": 2, "parallel": 2, "seed": 11}.items():
            batch.spin[key].setValue(value)
        batch.tau.setText("1")
        batch.output.setText(str(output))
        assert not batch.save_results.isChecked()
        before_files = set(output.rglob("*"))
        assert batch.generate() and len(batch.tasks) == 8
        paused = []
        states = []
        batch.runner.state_changed.connect(states.append)
        def pause_once(task):
            if task["progress"] > 0 and task["status"] == "running" and not paused:
                paused.append(True)
                batch.runner.pause()
                assert len(batch.runner.controllers) <= 2
                window.switch_workspace(0)
                window.switch_workspace(1)
                QTimer.singleShot(700, batch.runner.resume)
        batch.runner.task_changed.connect(pause_once)
        batch.start()
        wait(lambda: not batch.runner.active)
        batch.runner.task_changed.disconnect(pause_once)
        assert "paused" in states and states[-1] == "completed", states
        assert all(t["status"] == "completed" for t in batch.tasks), [(t["status"], t.get("error")) for t in batch.tasks]
        assert batch.stats.topLevelItemCount() == 4
        assert batch.runner.directory is None and batch.runner.report_path is None
        assert set(output.rglob("*")) == before_files
        for task in batch.tasks:
            assert "path" not in task["result"]
            restored = task["result"]["frames"]
            assert {frame["t"] for frame in restored} == {0, 1, 2}
            assert restored[0]["settings"]["search_algorithm_params"]["seed"] == task["seed"]
        for width, height in ((1360, 860), (900, 680), (700, 560)):
            window.resize(width, height)
            app.processEvents()
            assert window.width() <= width + 10, (width, window.width())
            assert window.grab().save(str(output / f"batch-{width}.png"))
        batch.replay_item(batch.task_table.topLevelItem(0))
        wait(lambda: not window.io_busy)
        assert window.workspace_stack.currentIndex() == 0
        assert window.frames
        assert window.chart_dashboard.visible_modes() == ["PF", "PS", "IGD"]
        window.resize(1360, 860)
        window.render_latest()
        app.processEvents()
        assert window.grab().save(str(output / "batch-result-replay.png"))
        print("PASS 8 parallel jobs / no disk writes by default / fair seeds / pause-resume / in-memory replay / 3 sizes", flush=True)
        batch.runner.export_statistics(output)
        wait(lambda: not batch.runner.active)
        book = load_workbook(batch.runner.report_path, read_only=True)
        assert book["逐次结果"].max_row == 9
        assert book["汇总统计"].max_row == 5
        book.close()
        assert not list(batch.runner.directory.glob("task_*.json"))
        print("PASS manual CSV-XLSX export without persisting raw snapshots", flush=True)

        runner = BatchRunner()
        tasks = copy.deepcopy(batch.tasks[:2])
        tasks[0]["request"]["params"]["search"]["unknown_batch_parameter"] = 1
        runner.start(tasks, 2, output, save_results=True)
        wait(lambda: not runner.active)
        assert [t["status"] for t in runner.tasks] == ["failed", "completed"]
        assert runner.report_path
        print("PASS one failed worker does not cancel the other job; failures excluded from statistics", flush=True)

        # Generate both kinds of component in an ignored temporary namespace,
        # prove they run in spawn workers, then remove only that fixture.
        import tempfile
        from pathlib import Path
        from flexdmo_app.components import create_scaffold, discover_plugins
        from flexdmo_app.experiments import build_plan
        from unittest.mock import patch
        with tempfile.TemporaryDirectory(prefix="plugin_test_", dir=ROOT / "results") as directory:
            create_scaffold("dynamic", "PreviewSmokeResponse", "Smoke response template", 2026, root=directory)
            create_scaffold("search", "PreviewSmokeSearch", "Smoke search template", 2026, root=directory)
            discovered, errors = discover_plugins(directory)
            assert not errors
            original_discover = discover_plugins
            with patch("flexdmo_app.components.discover_plugins", side_effect=lambda: original_discover(directory)):
                window.refresh_registry()
                assert any(r["name"] == "Smoke response template" for r in window.registry["dynamic"])
                assert any(batch.lists["search"].item(i).data(Qt.ItemDataRole.UserRole)["name"] == "Smoke search template"
                           for i in range(batch.lists["search"].count()))
            selection = {"dynamic": discovered["dynamic"], "search": discovered["search"],
                         "problem": [batch.selection()["problem"][0]]}
            shared = dict(batch.shared(), repeats=1)
            runner.start(build_plan(selection, shared), 1, output, save_results=True)
            wait(lambda: not runner.active)
            assert runner.tasks[0]["status"] == "completed", runner.tasks
            result = load_frames(runner.tasks[0]["result"]["path"])
            assert result[0]["settings"]["response_strategy_class"] == "PreviewSmokeResponse"
            assert result[0]["settings"]["search_algorithm_class"] == "PreviewSmokeSearch"
        window.refresh_registry()
        print("PASS generated dynamic/search components: auto-discovery in both UIs and real spawned run", flush=True)

        tasks = copy.deepcopy(batch.tasks[:3])
        for task in tasks:
            task["request"]["params"]["problem"]["total_evaluate_time"] = 10000
        runner.start(tasks, 1, output)
        # Waiting for the first controller message avoids unnecessarily running
        # hundreds of environments just to reach 1% of a very long plan.
        progress_seen = []
        runner.task_changed.connect(lambda task: progress_seen.append(True) if task.get("progress") == 0 and task["status"] == "running" else None)
        wait(lambda: bool(progress_seen) or runner.controllers and any(c.inbox and not c.inbox.empty() for c in runner.controllers.values()), timeout=30)
        runner.pause()
        runner.cancel()
        wait(lambda: not runner.active, timeout=15)
        assert all(t["status"] == "canceled" for t in runner.tasks)
        assert not runner.controllers
        assert runner.directory is None and runner.report_path is None
        print("PASS cancellation clears running and pending jobs without saving data", flush=True)
        runner.shutdown()
        # Configuration imports are transactional and component names portable.
        config = {"version": 1, "selection": {k: [v[0]["name"]] for k, v in batch.selection().items()},
                  "shared": batch.shared(), "profiles": []}
        batch.apply_configuration(config)
        before = batch.shared()
        invalid = copy.deepcopy(config)
        invalid["selection"]["dynamic"] = ["MissingAlgorithm"]
        try:
            batch.apply_configuration(invalid)
            raise AssertionError("invalid component accepted")
        except ValueError:
            assert batch.shared() == before
        window.dirty = False
        window.close()
        print("ALL BATCH INTEGRATION CHECKS PASSED", flush=True)
        return 0
    finally:
        window.controller.shutdown()
        batch.runner.shutdown()
        window.dirty = False
        window.close()


if __name__ == "__main__":
    raise SystemExit(main())
