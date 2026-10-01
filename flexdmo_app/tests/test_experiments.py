import copy
import json
from pathlib import Path
import tempfile
import unittest

from flexdmo_app.components import create_scaffold, discover_plugins
from flexdmo_app.core import records
from flexdmo_app.experiments import build_plan, export_reports, grouped_statistics, parse_sweep


class ExperimentTests(unittest.TestCase):
    def setUp(self):
        registry = records()
        self.selection = {kind: [next(r for r in rows if r["name"] == name)] for kind, rows, name in
                          (("dynamic", registry["dynamic"], "NoResponse"),
                           ("search", registry["search"], "NSGAII"), ("problem", registry["problem"], "CDP1"))}
        self.shared = {"tau": "1", "n": "10", "repeats": 2, "seed": 11,
                       "decision_num": 3, "solution_num": 6, "total_evaluate_time": 3}

    def test_sweep_commas_ranges_and_deduplication(self):
        self.assertEqual(parse_sweep("5，10,5:15:5"), [5, 10, 15])
        self.assertEqual(parse_sweep("1:3"), [1, 2, 3])

    def test_bad_sweep_is_rejected(self):
        for text in ("", "0", "1,", "1.5", "-1", "5:1", "1:5:0", "1:5000"):
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_sweep(text)

    def test_seed_schedule_is_fair_across_algorithms(self):
        self.selection["search"].append(next(r for r in records()["search"] if r["name"] == "SPEA2"))
        self.shared["tau"] = "1,2"
        tasks = build_plan(self.selection, self.shared)
        self.assertEqual(len(tasks), 8)
        for repeat, seed in ((1, 11), (2, 12)):
            self.assertEqual({t["seed"] for t in tasks if t["repeat"] == repeat}, {seed})
            self.assertEqual({t["request"]["params"]["search"]["seed"] for t in tasks if t["repeat"] == repeat}, {seed})

    def test_parameter_overrides_and_source_configs_untouched(self):
        record = self.selection["search"][0]
        path = Path(record["folder_name"]) / "config.json"
        before = path.read_bytes()
        tasks = build_plan(self.selection, self.shared, {("search", record["folder_name"]): {"proC": 0.75}})
        self.assertEqual(tasks[0]["request"]["params"]["search"]["proC"], 0.75)
        self.assertEqual(tasks[0]["request"]["params"]["problem"]["solution_num"], 6)
        self.assertEqual(path.read_bytes(), before)

    def test_empty_selection_and_seed_overflow_rejected(self):
        self.selection["dynamic"] = []
        with self.assertRaises(ValueError):
            build_plan(self.selection, self.shared)
        self.setUp()
        self.shared["seed"] = 2**32 - 1
        with self.assertRaises(ValueError):
            build_plan(self.selection, self.shared)

    def test_job_limit(self):
        self.shared.update(repeats=500, tau="1:21")
        with self.assertRaises(ValueError):
            build_plan(self.selection, self.shared)

    def test_group_mean_std_excludes_failed_and_partial(self):
        self.shared["repeats"] = 4
        tasks = build_plan(self.selection, self.shared)
        for task, score, status in zip(tasks, (0.1, 0.3, 0.0, 0.0), ("completed", "completed", "failed", "canceled")):
            task.update(status=status, result={"metrics": {"MIGD": score, "MGD": score, "MHV": 0.5, "feasibility": 1.0}})
        row = grouped_statistics(tasks)[0]
        self.assertEqual(row["完成次数"], 2)
        self.assertAlmostEqual(row["MIGD均值"], 0.2)
        self.assertAlmostEqual(row["MIGD标准差"], 0.2 / 2**0.5)

    def test_infinite_igd_is_not_silently_omitted(self):
        tasks = build_plan(self.selection, self.shared)
        for task, score in zip(tasks, (0.2, float("inf"))):
            task.update(status="completed", result={"metrics": {"MIGD": score}})
        row = grouped_statistics(tasks)[0]
        self.assertEqual(row["MIGD均值"], float("inf"))
        self.assertIsNone(row["MIGD标准差"])

    def test_csv_excel_manifest_keep_infinite_metric_explicit(self):
        from openpyxl import load_workbook
        tasks = build_plan(self.selection, self.shared)
        tasks[0].update(status="completed", result={"metrics": {"MIGD": float("inf"), "MHV": None}})
        with tempfile.TemporaryDirectory() as directory:
            path = export_reports(directory, tasks)
            book = load_workbook(path, read_only=True)
            self.assertEqual(book.sheetnames, ["逐次结果", "汇总统计"])
            self.assertEqual(book["逐次结果"].cell(2, 10).value, "inf")
            book.close()
            raw = (Path(directory) / "manifest.json").read_text()
            self.assertNotIn("Infinity", raw)
            self.assertEqual(json.loads(raw)[0]["result"]["MIGD"], "inf")
            self.assertTrue((Path(directory) / "summary.csv").exists())

    def test_scaffolds_are_discoverable_without_importing_code(self):
        with tempfile.TemporaryDirectory() as directory:
            for kind, name in (("dynamic", "TrialResponse"), ("search", "TrialSearch")):
                path = create_scaffold(kind, name, name + " template", 2026, root=directory)
                self.assertTrue((path / "main.py").is_file())
                compile((path / "main.py").read_text(), str(path / "main.py"), "exec")
            found, errors = discover_plugins(directory)
            self.assertFalse(errors)
            self.assertEqual(len(found["dynamic"]), 1)
            self.assertEqual(len(found["search"]), 1)

    def test_component_rejects_collision_traversal_and_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                create_scaffold("dynamic", "../Bad", "bad", 2026, root=directory)
            with self.assertRaises(ValueError):
                create_scaffold("dynamic", "NoResponse", "other", 2026, root=directory)
            path = create_scaffold("dynamic", "AnotherTrial", "trial", 2026, root=directory)
            before = (path / "main.py").read_bytes()
            with self.assertRaises(ValueError):
                create_scaffold("dynamic", "AnotherTrial", "trial", 2026, root=directory)
            self.assertEqual((path / "main.py").read_bytes(), before)

    def test_invalid_component_is_reported_not_imported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_scaffold("dynamic", "BadTrial", "bad trial", 2026, root=directory)
            (path / "main.py").write_text("raise RuntimeError('must not run at discovery')\nclass WrongName: pass\n")
            found, errors = discover_plugins(directory)
            self.assertFalse(found["dynamic"])
            self.assertEqual(len(errors), 1)


if __name__ == "__main__":
    unittest.main()
